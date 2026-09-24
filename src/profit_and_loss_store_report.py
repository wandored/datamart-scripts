import os
import re
import pandas as pd
import argparse
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from db_utils.dbconnect import DatabaseConnection


def get_arguments():
    parser = argparse.ArgumentParser(
        description="Import PnL data and write to spreadsheet for each location"
    )
    parser.add_argument(
        "-y",
        "--year",
        type=int,
        required=True,
        help="The year for which the data is being imported.",
    )
    parser.add_argument(
        "-p",
        "--period",
        type=int,
        required=True,
        help="The period for which the data is being imported.",
    )
    parser.add_argument(
        "-w",
        "--week",
        type=int,
        required=True,
        help="The week for which the data is being imported.",
    )
    args = parser.parse_args()
    return args


def get_pnl_data(year, period, week):
    with DatabaseConnection() as db:
        query = """
            SELECT *
            FROM reporting.pnl_detail
            WHERE year = %s AND period = %s AND week <= %s
        """
        db.execute(query, (year, period, week))
        pnl_data = db.fetchall()

        df = pd.DataFrame(
            [
                {
                    "week": row["week"],
                    "period": row["period"],
                    "year": row["year"],
                    "location": row["location"],
                    "account_category": row["account_category"],
                    "amount": row["amount"],
                    "pnl_report_order": row["pnl_report_order"],
                    "percent_of_category": row["percent_of_category"],
                }
                for row in pnl_data
            ]
        )
        return df


def create_location_report(
    location,
    current_year_data,
    prior_year_data,
    period,
    through_week,
    output_dir,
):

    # ------------------------------------------------------------------
    # Filter to this location
    # ------------------------------------------------------------------

    current = current_year_data[current_year_data["location"] == location].copy()

    prior = prior_year_data[prior_year_data["location"] == location].copy()

    # ------------------------------------------------------------------
    # Build the complete category list.
    #
    # Current year is authoritative when a category exists in both
    # years.  Prior-year values are used for categories that only exist
    # in the prior year.
    # ------------------------------------------------------------------

    current_info = current[
        [
            "account_category",
            "pnl_report_order",
            "percent_of_category",
        ]
    ].drop_duplicates("account_category")

    prior_info = prior[
        [
            "account_category",
            "pnl_report_order",
            "percent_of_category",
        ]
    ].drop_duplicates("account_category")

    category_info = current_info.merge(
        prior_info,
        on="account_category",
        how="outer",
        suffixes=("_current", "_prior"),
    )

    category_info["pnl_report_order"] = category_info[
        "pnl_report_order_current"
    ].fillna(category_info["pnl_report_order_prior"])

    category_info["percent_of_category"] = category_info[
        "percent_of_category_current"
    ].fillna(category_info["percent_of_category_prior"])

    category_info = category_info[
        [
            "account_category",
            "pnl_report_order",
            "percent_of_category",
        ]
    ].sort_values(
        ["pnl_report_order", "account_category"],
        na_position="last",
    )

    # ------------------------------------------------------------------
    # Helper to calculate weekly $ and %
    # ------------------------------------------------------------------

    def calculate_week(data, week):
        week_data = data[data["week"] == week]

        amounts = week_data.groupby("account_category")["amount"].sum().rename("amount")

        result = category_info[["account_category", "percent_of_category"]].copy()

        result["amount"] = result["account_category"].map(amounts).fillna(0)

        denominator = result["percent_of_category"].map(amounts)

        result[f"W{week} $"] = result["amount"]

        result[f"W{week} %"] = result["amount"] / denominator.replace(0, pd.NA)

        return result[
            [
                "account_category",
                f"W{week} $",
                f"W{week} %",
            ]
        ]

    # ------------------------------------------------------------------
    # Start with the category list
    # ------------------------------------------------------------------

    report = category_info[
        [
            "account_category",
            "pnl_report_order",
            "percent_of_category",
        ]
    ].copy()

    # ------------------------------------------------------------------
    # Weekly columns
    # ------------------------------------------------------------------

    for week in range(1, through_week + 1):
        current_week = calculate_week(current, week)

        report = report.merge(
            current_week,
            on="account_category",
            how="left",
        )

    # ------------------------------------------------------------------
    # Current-year PTD
    # ------------------------------------------------------------------

    current_ptd_amounts = current.groupby("account_category")["amount"].sum()

    report["PTD $"] = report["account_category"].map(current_ptd_amounts).fillna(0)

    current_ptd_denominator = report["percent_of_category"].map(current_ptd_amounts)

    report["PTD %"] = report["PTD $"] / current_ptd_denominator.replace(0, pd.NA)

    # ------------------------------------------------------------------
    # Prior-year PTD
    # ------------------------------------------------------------------

    prior_ptd_amounts = prior.groupby("account_category")["amount"].sum()

    report["PY PTD $"] = report["account_category"].map(prior_ptd_amounts).fillna(0)

    prior_ptd_denominator = report["percent_of_category"].map(prior_ptd_amounts)

    report["PY PTD %"] = report["PY PTD $"] / prior_ptd_denominator.replace(0, pd.NA)

    # ------------------------------------------------------------------
    # Variance
    # ------------------------------------------------------------------

    report["PTD Var $"] = report["PTD $"] - report["PY PTD $"]

    report["PTD Var %"] = report["PTD Var $"] / report["PY PTD $"].replace(0, pd.NA)

    # ------------------------------------------------------------------
    # Sort by P&L report order
    # ------------------------------------------------------------------

    report = report.sort_values(
        ["pnl_report_order", "account_category"],
        na_position="last",
    )

    # ------------------------------------------------------------------
    # Select final Excel columns
    # ------------------------------------------------------------------

    columns = ["account_category"]

    for week in range(1, through_week + 1):
        columns.extend(
            [
                f"W{week} $",
                f"W{week} %",
            ]
        )

    columns.extend(
        [
            "PTD $",
            "PTD %",
            "PY PTD $",
            "PY PTD %",
            "PTD Var $",
            "PTD Var %",
        ]
    )

    report = report[columns]

    # Rename first column
    report = report.rename(columns={"account_category": "Account Category"})

    # ------------------------------------------------------------------
    # Create workbook
    # ------------------------------------------------------------------

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "P&L"

    # ------------------------------------------------------------------
    # Title
    # ------------------------------------------------------------------

    worksheet["A1"] = location
    worksheet["A1"].font = Font(size=16, bold=True)

    worksheet["A2"] = f"Period {period} — Through Week {through_week}"
    worksheet["A2"].font = Font(size=12, bold=True)

    header_row = 4

    # ------------------------------------------------------------------
    # Headers
    # ------------------------------------------------------------------

    for column_number, column_name in enumerate(
        report.columns,
        start=1,
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_number,
            value=column_name,
        )

        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # ------------------------------------------------------------------
    # Data
    # ------------------------------------------------------------------

    # Excel cannot write pandas.NA; convert missing values to None
    report = report.astype(object).where(report.notna(), None)

    for row_number, row in enumerate(
        report.itertuples(index=False),
        start=header_row + 1,
    ):
        for column_number, value in enumerate(
            row,
            start=1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )

    # ------------------------------------------------------------------
    # Number formatting
    # ------------------------------------------------------------------

    currency_format = "#,##0.00;[Red](#,##0.00)"
    percent_format = "0.0%"

    for column_number, column_name in enumerate(
        report.columns,
        start=1,
    ):
        if column_name == "Account Category":
            continue

        number_format = percent_format if "%" in column_name else currency_format

        for row_number in range(
            header_row + 1,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row_number,
                column=column_number,
            ).number_format = number_format

    # ------------------------------------------------------------------
    # Bold total/subtotal rows
    # ------------------------------------------------------------------

    for row_number in range(
        header_row + 1,
        worksheet.max_row + 1,
    ):
        account_category = worksheet.cell(
            row=row_number,
            column=1,
        ).value

        if account_category and (
            account_category.startswith("Total ") or account_category.startswith("Net ")
        ):
            for column_number in range(
                1,
                worksheet.max_column + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).font = Font(bold=True)

    # ------------------------------------------------------------------
    # Column widths
    # ------------------------------------------------------------------

    worksheet.column_dimensions["A"].width = 35

    for column_number in range(2, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column_number)].width = 14

    # ------------------------------------------------------------------
    # Freeze panes
    # ------------------------------------------------------------------

    worksheet.freeze_panes = "B5"

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------

    safe_location = re.sub(
        r'[\\/*?:"<>|]',
        "_",
        location,
    )

    os.makedirs(output_dir, exist_ok=True)

    filename = os.path.join(
        output_dir,
        f"{safe_location}.xlsx",
    )

    workbook.save(filename)
    print(f"{filename} saved.")

    return filename


def format_workbook(filepath, location, year, period, week):
    workbook = load_workbook(filepath)
    worksheet = workbook["P&L"]

    # Title
    worksheet["A1"] = location
    worksheet["A1"].font = Font(
        bold=True,
        size=16,
    )

    worksheet["A2"] = f"Period {period}  |  Week {week}  |  {year}"
    worksheet["A2"].font = Font(
        bold=True,
        size=11,
    )

    # Header row
    header_row = 5

    for cell in worksheet[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Currency formatting
    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        min_col=2,
        max_col=2,
    ):
        for cell in row:
            cell.number_format = "#,##0.00;[Red](#,##0.00)"

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        min_col=4,
        max_col=4,
    ):
        for cell in row:
            cell.number_format = "#,##0.00;[Red](#,##0.00)"

    for row in worksheet.iter_rows(
        min_row=header_row + 1,
        min_col=6,
        max_col=6,
    ):
        for cell in row:
            cell.number_format = "#,##0.00;[Red](#,##0.00)"

    # Percentage formatting
    for column in (3, 5, 7):
        for row in range(
            header_row + 1,
            worksheet.max_row + 1,
        ):
            worksheet.cell(
                row=row,
                column=column,
            ).number_format = "0.0%"

    # Find total/subtotal rows.
    # We can refine this once we identify exactly which
    # account_category values should be bold.
    total_keywords = (
        "Total ",
        "Net Profit",
    )

    for row in range(header_row + 1, worksheet.max_row + 1):
        account = worksheet.cell(row, 1).value

        if account and (account.startswith(total_keywords) or account == "Net Profit"):
            for cell in worksheet[row]:
                cell.font = Font(bold=True)

    # Borders
    thin = Side(style="thin")

    for cell in worksheet[header_row]:
        cell.border = Border(bottom=thin)

    # Column widths
    worksheet.column_dimensions["A"].width = 32

    for column in range(2, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 16

    # Freeze the report header
    worksheet.freeze_panes = "B6"

    workbook.save(filepath)


def main():
    args = get_arguments()

    pnl_current = get_pnl_data(args.year, args.period, args.week)
    pnl_prior = get_pnl_data(args.year - 1, args.period, args.week)

    locations = sorted(
        set(pnl_current["location"].dropna().unique())
        | set(pnl_prior["location"].dropna().unique())
    )

    output_dir = Path("output/reports") / f"P{args.period}_{args.year}"

    for location in locations:
        create_location_report(
            location=location,
            current_year_data=pnl_current,
            prior_year_data=pnl_prior,
            period=args.period,
            through_week=args.week,
            output_dir=output_dir,
        )


if __name__ == "__main__":
    main()
